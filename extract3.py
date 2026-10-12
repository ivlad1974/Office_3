import openpyxl, json, warnings
warnings.filterwarnings('ignore')

wb = openpyxl.load_workbook('Реестр ЗАКАЗОВ 2.xlsm', data_only=True)
ws = wb['РЕЕСТР']
hdr = [ws.cell(3, c).value for c in range(1, 20)]
rows = []
for r in range(4, 1000):
    vals = [ws.cell(r, c).value for c in range(1, 20)]
    if any(vals):
        rows.append(vals)

with open('public/data/db_rz.json', 'w', encoding='utf-8') as f:
    json.dump({'headers': hdr, 'rows': rows}, f, ensure_ascii=False, indent=2, default=str)
print('Extract3 complete:', len(rows), 'rows')
