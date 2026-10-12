import openpyxl, json, warnings
warnings.filterwarnings('ignore')

wb = openpyxl.load_workbook('Реестр 2.xlsm', data_only=True)
ws = wb['РасчетДокументация']
hdr = [ws.cell(10, c).value for c in range(1, 60)]
rows = []
for r in range(11, 900):
    vals = [ws.cell(r, c).value for c in range(1, 60)]
    if any(vals):
        rows.append(vals)

with open('public/data/db_reestr2.json', 'w', encoding='utf-8') as f:
    json.dump({'headers': hdr, 'rows': rows}, f, ensure_ascii=False, indent=2, default=str)
print('Extract2 complete:', len(rows), 'rows')
