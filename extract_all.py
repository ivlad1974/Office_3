import openpyxl, json, re, warnings
warnings.filterwarnings('ignore')

def cellval(ws, r, c):
    v = ws.cell(row=r, column=c).value
    if v is None: return None
    if isinstance(v, str) and v.startswith('='): return None
    return v

out = {}

# ---------- ЗАКАЗ ----------
try:
    wb = openpyxl.load_workbook('Заказ_3.1 1-Этап ДЗЕН от 11.08.2026.xlsm', data_only=False)
    m = wb['Материалы']
    meta = {
        'objectName': m.cell(1, 5).value, 'org': m.cell(2, 5).value, 'customer': m.cell(3, 5).value,
        'version': m.cell(1, 12).value, 'coefMat': m.cell(5, 9).value or 1, 'coefWork': m.cell(5, 10).value or 1,
    }
    rows = []
    for r in range(8, 2003):
        b = m.cell(r, 2).value; d = m.cell(r, 4).value
        f = m.cell(r, 6).value; g = m.cell(r, 7).value; e = m.cell(r, 5).value; k = m.cell(r, 11).value
        if d is None and f is None: continue
        if isinstance(d, str) and d.startswith('='): d = None
        typ = b if isinstance(b, str) else None
        qty = f if isinstance(f, (int, float)) else None
        price = g if isinstance(g, (int, float)) else None
        rows.append({'t': typ, 'n': d if isinstance(d, str) else None, 'u': e if isinstance(e, str) else None,
                     'q': qty, 'p': price, 'note': k if isinstance(k, str) else None})
    out['order_meta'] = meta
    out['order_rows'] = rows

    ks6 = wb['КС-6']
    k6 = {'stroyka': ks6.cell(7, 3).value, 'obekt': ks6.cell(9, 3).value,
          'dogovorNum': ks6.cell(13, 13).value,
          'dogovorDate': str(ks6.cell(14, 13).value)[:10] if ks6.cell(14, 13).value else None,
          'dates': [], 'status': []}
    for c in range(6, 15):
        v = ks6.cell(32, c).value
        k6['dates'].append(str(v)[:10] if v else None)
        st = ks6.cell(30, c).value
        k6['status'].append(st if isinstance(st, str) else '----')
    out['ks6_header'] = k6
except Exception as ex:
    print('Order error:', ex)

# ---------- КОНТРАГЕНТЫ ----------
try:
    with open('kontr.json', 'r', encoding='utf-8') as kf:
        out['kontragenty'] = json.load(kf)
except Exception:
    pass

# ---------- РЕЕСТР 2 ----------
try:
    wb2 = openpyxl.load_workbook('Реестр 2.xlsm', data_only=True)
    ws = wb2['РасчетДокументация']
    hdr = [ws.cell(10, c).value for c in range(1, 60)]
    orders = []
    for r in range(11, 900):
        vals = [ws.cell(r, c).value for c in range(1, 60)]
        if not any(vals): continue
        orders.append(vals)
    out['reestr_headers'] = hdr
    out['reestr_orders'] = orders
except Exception as ex:
    print('Reestr2 error:', ex)

# ---------- РЕЕСТР ЗАКАЗОВ 2 ----------
try:
    wb3 = openpyxl.load_workbook('Реестр ЗАКАЗОВ 2.xlsm', data_only=True)
    ws3 = wb3['РЕЕСТР']
    hdr3 = [ws3.cell(3, c).value for c in range(1, 20)]
    rz_orders = []
    for r in range(4, 1000):
        vals = [ws3.cell(r, c).value for c in range(1, 20)]
        if not any(vals): continue
        rz_orders.append(vals)
    out['rz_headers'] = hdr3
    out['rz_orders'] = rz_orders
except Exception as ex:
    print('RZ error:', ex)

with open('public/data/db.json', 'w', encoding='utf-8') as f:
    json.dump(out, f, ensure_ascii=False, indent=2, default=str)
print('Data extracted successfully')
